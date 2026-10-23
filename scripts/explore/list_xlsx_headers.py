"""Print sheet names, row counts and header rows of every .xlsx file in a folder.

Usage:
    python scripts/explore/list_xlsx_headers.py [folder]
"""
import sys
from pathlib import Path

import openpyxl

DEFAULT_FOLDER = "data/raw/aer_st37/ST_37_Excel"


def main(folder):
    for path in sorted(Path(folder).glob("*.xlsx")):
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb.worksheets[0]
        header = next(ws.iter_rows(max_row=1, values_only=True))
        print(f"{path.name} | sheets: {wb.sheetnames} | rows incl. header: {ws.max_row:,}")
        print(f"    {list(header)}")
        wb.close()


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_FOLDER)
